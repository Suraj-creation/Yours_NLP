"""Module 1: load any supported format into one Document type (Lab 1 pattern).

detect format -> pick the loader -> extract text -> return Document(id, text, metadata)

Formats: XML (OpenStax CNXML with MathML, ACL Anthology, generic), PDF, Markdown,
JSON, JSONL, CSV, XLSX, TXT, DOCX, HTML. Structured records are turned into
readable paragraphs ("records to paragraphs", Lab 1), and annotation codes
(dialogue acts, talk-move tags) are moved to metadata so they do not pollute
the language.
"""
from __future__ import annotations

import csv
import io
import json
import re
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from pathlib import Path

from .config import DATA, DATA_SCALE, RESULTS


@dataclass
class Document:
    doc_id: str
    filename: str
    fmt: str
    title: str
    text: str
    meta: dict = field(default_factory=dict)

    def to_dict(self, preview: int | None = None) -> dict:
        t = self.text if preview is None else self.text[:preview]
        return {"doc_id": self.doc_id, "filename": self.filename, "format": self.fmt,
                "title": self.title, "text": t, "meta": self.meta}


# ------------------------------------------------------------------ MathML
def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _join(parts: list[str]) -> str:
    out = ""
    for p in parts:
        if not p:
            continue
        if out and (out[-1].isalnum() and p[0].isalnum()):
            out += " "  # "2" + "1/4" -> "2 1/4" (mixed number); "x" "y" -> "x y"
        out += p
    return out


def _wrap(s: str) -> str:
    return f"({s})" if re.search(r"[+\-−×·=\s/]", s) else s


def mathml_to_text(el: ET.Element) -> str:
    """Linearise MathML: <mfrac>3 4</mfrac> -> '3/4', <msup>x 2</msup> -> 'x^2'."""
    tag = _local(el.tag)
    kids = list(el)
    if tag in ("mn", "mi", "mo", "mtext", "ms"):
        txt = (el.text or "").strip()
        return {"−": "-", "×": "×", "⋅": "·", "⁢": ""}.get(txt, txt)
    if tag == "mspace":
        return " "
    if tag in ("annotation", "annotation-xml"):
        return ""
    if tag == "mfrac" and len(kids) == 2:
        return f"{_wrap(mathml_to_text(kids[0]))}/{_wrap(mathml_to_text(kids[1]))}"
    if tag == "msup" and len(kids) == 2:
        return f"{mathml_to_text(kids[0])}^{_wrap(mathml_to_text(kids[1]))}"
    if tag == "msub" and len(kids) == 2:
        return f"{mathml_to_text(kids[0])}_{mathml_to_text(kids[1])}"
    if tag == "msubsup" and len(kids) == 3:
        return f"{mathml_to_text(kids[0])}_{mathml_to_text(kids[1])}^{mathml_to_text(kids[2])}"
    if tag == "msqrt":
        return f"sqrt({_join([mathml_to_text(k) for k in kids])})"
    if tag == "mroot" and len(kids) == 2:
        return f"root({mathml_to_text(kids[0])},{mathml_to_text(kids[1])})"
    if tag == "mfenced":
        return "(" + ",".join(mathml_to_text(k) for k in kids) + ")"
    if tag in ("mtable", "mtr"):
        return " ".join(mathml_to_text(k) for k in kids)
    return _join([mathml_to_text(k) for k in kids])


# ------------------------------------------------------------------ CNXML
_BLOCK = {"title", "para", "item", "entry", "caption", "problem", "solution", "label", "term-list",
          "definition", "meaning", "example", "note", "list", "table", "row", "section", "exercise",
          "equation", "figure"}
_SKIP = {"metadata", "link-target"}


def _cnxml_text(el: ET.Element, out: list[str], stats: dict) -> None:
    tag = _local(el.tag)
    if tag in _SKIP:
        return
    if tag == "math":
        stats["math"] = stats.get("math", 0) + 1
        out.append(mathml_to_text(el))
        if el.tail:
            out.append(el.tail)
        return
    if tag == "media":
        # Figure alt texts describe images ("A white background with the fraction 3/4 ...").
        # They are accessibility descriptions, not maths content, so they are counted and
        # left out by default (include_alt=True keeps them).
        alt = el.get("alt")
        if alt:
            stats["alt"] = stats.get("alt", 0) + 1
            if stats.get("_include_alt"):
                out.append("\n" + " ".join(alt.split()) + "\n")
        if el.tail:
            out.append(el.tail)
        return
    if tag == "newline":
        out.append("\n")
    if tag in _BLOCK:
        out.append("\n")
    if el.text:
        out.append(el.text)
    for k in el:
        _cnxml_text(k, out, stats)
    if tag in _BLOCK:
        out.append("\n")
    if el.tail:
        out.append(el.tail)


def load_cnxml(path: Path, include_alt: bool = False) -> tuple[str, str, dict]:
    root = ET.parse(path).getroot()
    ns = {"c": "http://cnx.rice.edu/cnxml", "md": "http://cnx.rice.edu/mdml"}
    title = (root.findtext("c:title", default="", namespaces=ns) or "").strip()
    stats: dict = {"_include_alt": include_alt}
    parts: list[str] = [title, "\n"]
    abstract = root.find(".//md:abstract", ns)
    if abstract is not None:
        _cnxml_text(abstract, parts, stats)
    content = root.find("c:content", ns)
    if content is not None:
        _cnxml_text(content, parts, stats)
    text = "".join(parts)
    return title, text, {"math_expressions": stats.get("math", 0), "alt_texts_removed": stats.get("alt", 0),
                         "module": root.findtext(".//md:content-id", default="", namespaces=ns)}


def load_acl_collection(path: Path) -> tuple[str, str, dict]:
    root = ET.parse(path).getroot()
    lines, n = [], 0
    for vol in root.iter("volume"):
        for p in vol.iter("paper"):
            title = " ".join("".join(p.find("title").itertext()).split())
            authors = []
            for a in p.findall("author"):
                name = " ".join(filter(None, [a.findtext("first"), a.findtext("last")]))
                aff = a.findtext("affiliation")
                authors.append(f"{name} ({aff})" if aff else name)
            abstract = " ".join("".join(p.find("abstract").itertext()).split()) if p.find("abstract") is not None else ""
            lines.append(f"{title}\nAuthors: {'; '.join(authors)}\nVenue: {vol.get('id')}\n{abstract}\n")
            n += 1
    return root.get("title", path.stem), "\n".join(lines), {"papers": n}


def load_xml(path: Path) -> tuple[str, str, dict, str]:
    head = path.read_text(encoding="utf-8", errors="ignore")[:500]
    if "cnx.rice.edu/cnxml" in head:
        return (*load_cnxml(path), "cnxml+mathml")
    if 'source="ACL Anthology"' in head:
        return (*load_acl_collection(path), "acl-anthology-xml")
    root = ET.parse(path).getroot()
    return path.stem, " ".join(root.itertext()), {}, "xml-itertext"


# ------------------------------------------------------------------ PDF
# Default is pdfplumber: on D13 it keeps 83% dictionary words vs 69% for pypdf,
# which breaks words apart ("T alk Mo ves"). See compare_pdf_extractors().
def pdf_pages(path: Path, method: str = "pdfplumber") -> list[str]:
    if method == "pypdf":
        from pypdf import PdfReader

        return [p.extract_text() or "" for p in PdfReader(str(path)).pages]
    if method == "pdfplumber":
        import pdfplumber

        with pdfplumber.open(str(path)) as pdf:
            return [p.extract_text() or "" for p in pdf.pages]
    raise ValueError(method)


def compare_pdf_extractors(path: Path) -> list[dict]:
    """Lab 1, Ex 4 style comparison: which extractor gives cleaner text?"""
    from nltk.corpus import words as nltk_words

    english = {w.lower() for w in nltk_words.words()}
    rows = []
    for m in ("pypdf", "pdfplumber"):
        import time

        t0 = time.perf_counter()
        pages = pdf_pages(path, m)
        dt = time.perf_counter() - t0
        text = "\n".join(pages)
        toks = re.findall(r"[A-Za-z]+", text)
        known = sum(1 for t in toks if t.lower() in english)
        rows.append({"method": m, "pages": len(pages), "characters": len(text), "words": len(toks),
                     "dictionary_word_share": round(known / max(1, len(toks)), 4),
                     "hyphen_breaks": len(re.findall(r"[a-z]-\n[a-z]", text)),
                     "seconds": round(dt, 3), "sample": text[:300]})
    return rows


def load_pdf(path: Path, method: str = "pdfplumber") -> tuple[str, str, dict]:
    pages = pdf_pages(path, method)
    title = next((ln.strip() for ln in pages[0].splitlines() if ln.strip()), path.stem) if pages else path.stem
    return title, "\n\n".join(pages), {"pages": len(pages), "extractor": method}


# ------------------------------------------------------------------ Markdown / TXT / DOCX / HTML
def load_markdown(path: Path) -> tuple[str, str, dict]:
    md = path.read_text(encoding="utf-8")
    title = next((ln.lstrip("# ").strip() for ln in md.splitlines() if ln.startswith("#")), path.stem)
    t = re.sub(r"```.*?```", " ", md, flags=re.S)             # code blocks are not language
    t = re.sub(r"<!--.*?-->", " ", t, flags=re.S)
    t = re.sub(r"!\[[^\]]*\]\([^)]*\)", " ", t)               # images and badges
    t = re.sub(r"\[!\[[^\]]*\]\[[^\]]*\]\]\[[^\]]*\]", " ", t)
    t = re.sub(r"^\[[^\]]+\]:\s*\S+.*$", " ", t, flags=re.M)  # reference-style link targets
    t = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", t)            # [text](url) -> text
    t = re.sub(r"\[([^\]]+)\]\[[^\]]*\]", r"\1", t)
    t = re.sub(r"<[^>]+>", " ", t)
    t = re.sub(r"^#+\s*", "", t, flags=re.M)
    t = re.sub(r"`([^`]*)`", r"\1", t)
    t = re.sub(r"(\*\*|__|\*|_)(\S[^*_]*?\S?)\1", r"\2", t)
    t = re.sub(r"^\s*[-*+]\s+", "", t, flags=re.M)
    t = t.replace("|", " ")
    return title, t, {}


def load_txt(path: Path) -> tuple[str, str, dict]:
    text = path.read_text(encoding="utf-8")
    return path.stem, text, {"lines": len(text.splitlines())}


def load_docx(path: Path) -> tuple[str, str, dict]:
    from docx import Document as Docx

    d = Docx(str(path))
    paras = [p.text for p in d.paragraphs if p.text.strip()]
    rows = []
    for tb in d.tables:
        for r in tb.rows[1:]:
            cells = [c.text.strip() for c in r.cells]
            rows.append(". ".join(c.rstrip(".") for c in cells if c) + ".")
    title = paras[0] if paras else path.stem
    # headings first, then the observation table, then the remaining paragraphs (reading order)
    text = "\n".join(paras[:6] + rows + paras[6:])
    return title, text, {"paragraphs": len(paras), "table_rows": len(rows)}


def html_extract(html: str, method: str = "trafilatura") -> str:
    if method == "trafilatura":
        import trafilatura

        return trafilatura.extract(html, include_tables=True, include_comments=False, favor_recall=True) or ""
    from bs4 import BeautifulSoup

    return BeautifulSoup(html, "html.parser").get_text(separator="\n")


def compare_html_extractors(path: Path) -> list[dict]:
    html = path.read_text(encoding="utf-8")
    boiler = ["cookies", "Accept all", "Admissions", "Parent Login", "Privacy policy", "All rights reserved",
              "social media", "tracking", "window.analytics"]
    rows = []
    for m in ("trafilatura", "beautifulsoup"):
        t = html_extract(html, m)
        rows.append({"method": m, "characters": len(t), "words": len(t.split()),
                     "boilerplate_phrases_found": sum(1 for b in boiler if b.lower() in t.lower()),
                     "sample": " ".join(t.split())[:300]})
    return rows


def load_html(path: Path) -> tuple[str, str, dict]:
    from bs4 import BeautifulSoup

    html = path.read_text(encoding="utf-8")
    title = BeautifulSoup(html, "html.parser").title
    return (title.get_text(strip=True) if title else path.stem), html_extract(html, "trafilatura"), {"extractor": "trafilatura"}


# ------------------------------------------------------------------ JSON / JSONL / CSV / XLSX
_ACT = re.compile(r"^\((\w+)\)\s*")


def mathdial_record_to_text(r: dict, with_conversation: bool = True) -> tuple[str, dict]:
    # Field labels are kept short and neutral on purpose: a label such as
    # "Teacher-described confusion" repeated in every record would make every document
    # match "teacher" queries because of the conversion, not the content.
    acts: dict = {}
    lines = [f"Problem: {r.get('question', '').strip()}"]
    if r.get("student_profile"):
        lines.append(f"Profile: {r['student_profile'].strip()}")
    if r.get("student_incorrect_solution"):
        lines.append(f"Incorrect solution: {r['student_incorrect_solution'].strip()}")
    if r.get("teacher_described_confusion"):
        lines.append(f"Confusion: {r['teacher_described_confusion'].strip()}")
    if with_conversation and r.get("conversation"):
        lines.append("Conversation:")
        for turn in r["conversation"].split("|EOM|"):
            turn = turn.strip()
            if not turn:
                continue
            spk, _, utt = turn.partition(":")
            utt = utt.strip()
            m = _ACT.match(utt)
            if m:
                acts[m.group(1)] = acts.get(m.group(1), 0) + 1
                utt = utt[m.end():]
            lines.append(f"{spk.strip()}: {utt}")
    return "\n".join(lines), acts


def load_jsonl(path: Path) -> tuple[str, str, dict]:
    records = [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]
    blocks, acts = [], {}
    for r in records:
        t, a = mathdial_record_to_text(r)
        for k, v in a.items():
            acts[k] = acts.get(k, 0) + v
        blocks.append(t)
    return path.stem, "\n\n".join(blocks), {"records": len(records), "dialogue_acts_removed": acts}


def load_json(path: Path) -> tuple[str, str, dict]:
    data = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(data, list) and data and "Misconception ID" in data[0]:  # MaE
        blocks = []
        for r in data:
            b = [f"Misconception {r['Misconception ID']} ({r['Topic']}): {r['Misconception'].strip()}",
                 f"Question: {' '.join(str(r.get('Question', '')).split())}",
                 f"Incorrect answer: {r.get('Incorrect Answer', '')}",
                 f"Correct answer: {r.get('Correct Answer', '')}"]
            if r.get("Explanation"):
                b.append(f"Explanation: {r['Explanation']}")
            if r.get("Source"):
                b.append(f"Source: {r['Source']}")
            blocks.append("\n".join(b))
        return "MaE misconceptions", "\n\n".join(blocks), {"records": len(data)}
    if isinstance(data, dict) and "sessions" in data:  # tutor chat log
        lines = [f"Learner: {data.get('learner_name')} ({data.get('learner_id')}). Tutor: {data.get('tutor')}. "
                 f"Platform: {data.get('platform')}."]
        for s in data["sessions"]:
            lines.append(f"Session {s['session']} on {s['date']}")
            for t in s["turns"]:
                who = data.get("tutor") if t["speaker"] == "tutor" else data.get("learner_name")
                lines.append(f"[{t['t']}] {who}: {t['text']}")
        return "Tutor chat log", "\n".join(lines), {"sessions": len(data["sessions"])}

    def walk(x, prefix=""):
        if isinstance(x, dict):
            for k, v in x.items():
                yield from walk(v, f"{k}")
        elif isinstance(x, list):
            for v in x:
                yield from walk(v, prefix)
        elif isinstance(x, str):
            yield f"{prefix}: {x}" if prefix else x

    return path.stem, "\n".join(walk(data)), {}


def load_csv(path: Path) -> tuple[str, str, dict]:
    with open(path, newline="", encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    blocks = []
    for r in rows:
        if "student_incorrect_solution" in r:
            blocks.append("\n".join([
                f"Problem: {r['question'].strip()}",
                f"Incorrect solution: {r['student_incorrect_solution'].strip()}",
                f"Confusion: {r['teacher_described_confusion'].strip()}",
                f"Solved: {r.get('self-correctness', '')}"]))
        else:
            blocks.append("\n".join(f"{k}: {v}" for k, v in r.items() if v))
    return path.stem, "\n\n".join(blocks), {"records": len(rows)}


_TEACHER = re.compile(r"^(T|T\d|T/R\d|R\d|Teacher)$", re.I)


def talkmoves_rows_to_text(df) -> tuple[str, dict]:
    lines, tags = [], {}
    tag_cols = [c for c in df.columns if "Tag" in str(c)]
    for _, r in df.iterrows():
        s = r.get("Sentence")
        if not isinstance(s, str) or not s.strip():
            continue
        spk = str(r.get("Speaker", "")).strip()
        spk = "Teacher" if _TEACHER.match(spk) else ("Student" if spk in ("S", "nan", "") else spk)
        lines.append(f"{spk}: {' '.join(s.split())}")
        for c in tag_cols:
            v = r.get(c)
            if isinstance(v, str) and v.strip():
                tags[v.strip()] = tags.get(v.strip(), 0) + 1
    return "\n".join(lines), tags


def load_xlsx(path: Path) -> tuple[str, str, dict]:
    import pandas as pd

    df = pd.read_excel(path)
    text, tags = talkmoves_rows_to_text(df)
    return path.stem, text, {"rows": len(df), "talk_move_tags_removed": tags}


# ------------------------------------------------------------------ dispatch
LOADERS = {".xml": "xml", ".cnxml": "xml", ".pdf": "pdf", ".md": "markdown", ".txt": "txt", ".docx": "docx",
           ".html": "html", ".htm": "html", ".json": "json", ".jsonl": "jsonl", ".csv": "csv", ".xlsx": "xlsx"}


def load_document(path: Path, doc_id: str | None = None, title: str | None = None) -> Document:
    path = Path(path)
    kind = LOADERS.get(path.suffix.lower())
    if kind is None:
        raise ValueError(f"unsupported format: {path.suffix}")
    extractor = kind
    if kind == "xml":
        t, text, meta, extractor = load_xml(path)
    else:
        t, text, meta = {"pdf": load_pdf, "markdown": load_markdown, "txt": load_txt, "docx": load_docx,
                         "html": load_html, "json": load_json, "jsonl": load_jsonl, "csv": load_csv,
                         "xlsx": load_xlsx}[kind](path)
    meta = {**meta, "loader": extractor, "bytes": path.stat().st_size}
    did = doc_id or (path.name.split("_", 1)[0] if re.match(r"D\d\d_", path.name) else path.stem)
    return Document(did, path.name, kind.upper() if kind != "markdown" else "MD", title or t, text, meta)


def read_catalog() -> list[dict]:
    p = RESULTS / "document_catalog.csv"
    if not p.exists():
        return []
    with open(p, newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def load_corpus(data_dir: Path = DATA) -> list[Document]:
    cat = {c["doc_id"]: c for c in read_catalog()}
    docs = []
    for p in sorted(Path(data_dir).glob("D[0-9][0-9]_*")):
        did = p.name.split("_", 1)[0]
        c = cat.get(did, {})
        d = load_document(p, did, c.get("title"))
        d.meta.update({k: c[k] for k in ("layer", "source", "licence", "synthetic") if k in c})
        docs.append(d)
    return docs


# ------------------------------------------------------------------ heavy datasets (scale tier)
def load_mathdial_scale() -> list[Document]:
    docs = []
    for split in ("train", "test"):
        for line in open(DATA_SCALE / "mathdial" / f"{split}.jsonl", encoding="utf-8"):
            r = json.loads(line)
            text, acts = mathdial_record_to_text(r)
            i = len(docs) + 1
            docs.append(Document(f"MD{i:05d}", f"{split}.jsonl", "JSONL",
                                 " ".join(r["question"].split())[:90], text,
                                 {"split": split, "qid": r["qid"], "dialogue_acts": acts,
                                  "self_correctness": r.get("self-correctness"), "collection": "mathdial"}))
    return docs


def load_talkmoves_scale() -> list[Document]:
    import pandas as pd

    docs = []
    for f, split in (("train_data_504.xlsx", "train"), ("test_data_63.xlsx", "test")):
        df = pd.read_excel(DATA_SCALE / "talkmoves" / f)
        for name, g in df.groupby("Transcript", sort=False):
            text, tags = talkmoves_rows_to_text(g)
            i = len(docs) + 1
            docs.append(Document(f"TM{i:04d}", f, "XLSX", str(name).replace(".xlsx", ""), text,
                                 {"split": split, "rows": len(g), "collection": "talkmoves"}))
    return docs


def load_bytes(name: str, content: bytes) -> Document:
    """Load an uploaded file (website upload) through the same loaders."""
    import tempfile

    suffix = Path(name).suffix
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as fh:
        fh.write(content)
        tmp = Path(fh.name)
    try:
        d = load_document(tmp, doc_id="UPLOAD", title=Path(name).stem)
        d.filename = name
        return d
    finally:
        tmp.unlink(missing_ok=True)


__all__ = ["Document", "load_document", "load_corpus", "load_mathdial_scale", "load_talkmoves_scale",
           "compare_pdf_extractors", "compare_html_extractors", "mathml_to_text", "read_catalog", "load_bytes",
           "io"]
