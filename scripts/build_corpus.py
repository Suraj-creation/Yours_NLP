"""Build the 30-document assessment corpus (data/) and the heavy scale tier (data_scale/).

    python scripts/fetch_data.py      # once
    python scripts/build_corpus.py

Assessment corpus D01-D30: the documents every exercise, table and qrels use.
Scale tier: the two heavy datasets in full (MathDial, TalkMoves), processed by
the same pipeline to show the system at realistic volume.
"""
from __future__ import annotations

import csv
import json
import random
import re
import shutil
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import synthetic_content as S  # noqa: E402

SRC, DATA, SCALE, GOLD = ROOT / "sources", ROOT / "data", ROOT / "data_scale", ROOT / "gold"
random.seed(42)

LIC = {
    "openstax": "CC BY-NC-SA 4.0",
    "acl": "CC BY 4.0 (ACL Anthology metadata)",
    "mathdial": "CC BY-SA 4.0",
    "talkmoves": "CC BY-NC-SA 4.0",
    "mae": "MIT",
    "synthetic": "Created for this project (synthetic, labelled)",
}
URL = {
    "openstax": "https://github.com/openstax/osbooks-prealgebra-bundle",
    "acl": "https://github.com/acl-org/acl-anthology",
    "mathdial": "https://github.com/eth-nlped/mathdial",
    "talkmoves": "https://github.com/SumnerLab/TalkMoves",
    "mae": "https://github.com/nancyotero-projects/math-misconceptions",
}

catalog: list[dict] = []


def add(doc_id, filename, title, fmt, layer, source, synthetic=False, notes=""):
    catalog.append({
        "doc_id": doc_id, "filename": filename, "title": title, "format": fmt,
        "layer": layer, "source": source, "source_url": URL.get(source, ""),
        "licence": LIC["synthetic" if synthetic else source], "synthetic": synthetic, "notes": notes,
    })


# ---------------------------------------------------------------- OpenStax (XML)
def build_openstax(start: int) -> int:
    modules = json.loads((SRC / "manifest.json").read_text())["openstax_modules"]
    n = start
    for mid, title in modules.items():
        slug = re.sub(r"[^a-z0-9]+", "_", title.lower()).strip("_")
        fname = f"D{n:02d}_openstax_{slug}.xml"
        shutil.copy(SRC / "openstax" / f"{mid}.cnxml", DATA / fname)
        add(f"D{n:02d}", fname, f"OpenStax Prealgebra 2e, section {title}", "XML (CNXML + MathML)",
            "concept", "openstax", notes=f"module {mid}")
        n += 1
    return n


# ---------------------------------------------------------------- ACL Anthology (XML)
STRONG = re.compile(r"knowledge tracing|misconception|tutor(?:ing|s)?\b|classroom|pedagog|educational|"
                    r"student(?:s|s')? (?:answer|response|essay|question|explanation|misconception|"
                    r"knowledge|performance|writing|learning)|math(?:ematics)? (?:education|word problems?)|"
                    r"feedback on student|intelligent tutoring", re.I)
WEAK = re.compile(r"\bstudents?\b|\bteachers?\b|\beducation\b|\blearners?\b|\bschool\b|\bmath", re.I)
NOISE = re.compile(r"teacher[- ]student|student model|knowledge distillation|distil|in-context learners?", re.I)


def _text(el):
    return "".join(el.itertext()) if el is not None else ""


def select_papers(volume: str, must: str | None = None, cap: int = 30) -> list[ET.Element]:
    tree = ET.parse(SRC / "acl" / f"{volume}.xml")
    scored = []
    for p in tree.iter("paper"):
        title, abstract = _text(p.find("title")), _text(p.find("abstract"))
        if not abstract:
            continue
        blob = f"{title} {abstract}"
        strong, weak = len(STRONG.findall(blob)), len(WEAK.findall(blob))
        score = 2 * strong + weak
        if must and re.search(must, title, re.I):
            score += 100
        if strong == 0 or score < 3 or (NOISE.search(blob) and strong < 2):
            continue
        scored.append((score, p))
    scored.sort(key=lambda x: -x[0])
    return [p for _, p in scored[:cap]]


def write_collection(fname: str, parts: list[tuple[str, list[ET.Element]]], title: str) -> int:
    root = ET.Element("collection", {"source": "ACL Anthology", "title": title})
    count = 0
    for vol, papers in parts:
        v = ET.SubElement(root, "volume", {"id": vol})
        for p in papers:
            v.append(p)
            count += 1
    ET.indent(root)
    ET.ElementTree(root).write(DATA / fname, encoding="utf-8", xml_declaration=True)
    return count


def build_acl(n: int) -> int:
    specs = [
        ("acl_2025_education_kt", "ACL 2025: knowledge tracing, tutoring and education papers (incl. SQKT)",
         [("2025.acl", select_papers("2025.acl", must=r"Knowledge Tracing in Programming", cap=30))]),
        ("emnlp_2025_education_kt", "EMNLP 2025: knowledge tracing and education papers (incl. CIKT)",
         [("2025.emnlp", select_papers("2025.emnlp", must=r"CIKT", cap=30))]),
        ("bea_2023_2025_feedback_misconceptions", "BEA workshops 2023-2025: feedback, tutoring and misconception papers",
         [(v, select_papers(v, cap=12)) for v in ["2023.bea", "2024.bea", "2025.bea"]]),
        ("dataset_papers_mathdial_talkmoves_mae", "Dataset papers: MathDial, TalkMoves, MaE and related education datasets",
         [("2023.findings", select_papers("2023.findings", must=r"Math.*Dial", cap=10)),
          ("2022.lrec", select_papers("2022.lrec", must=r"Talk.*Moves", cap=10)),
          ("2025.aimecon", select_papers("2025.aimecon", must=r"Misconception", cap=10))]),
    ]
    for slug, title, parts in specs:
        fname = f"D{n:02d}_{slug}.xml"
        k = write_collection(fname, parts, title)
        add(f"D{n:02d}", fname, title, "XML (ACL Anthology)", "research", "acl", notes=f"{k} papers")
        n += 1
    return n


# ---------------------------------------------------------------- dataset docs
def build_dataset_docs(n: int) -> int:
    shutil.copy(SRC / "talkmoves" / "Coding Manual.pdf", DATA / f"D{n:02d}_talkmoves_coding_manual.pdf")
    add(f"D{n:02d}", f"D{n:02d}_talkmoves_coding_manual.pdf", "TalkMoves coding manual (teacher and student talk moves)",
        "PDF", "research", "talkmoves"); n += 1
    shutil.copy(SRC / "mathdial" / "README.md", DATA / f"D{n:02d}_mathdial_readme.md")
    add(f"D{n:02d}", f"D{n:02d}_mathdial_readme.md", "MathDial dataset card (README)", "Markdown", "research", "mathdial"); n += 1
    shutil.copy(SRC / "mae" / "README.md", DATA / f"D{n:02d}_mae_readme.md")
    add(f"D{n:02d}", f"D{n:02d}_mae_readme.md", "MaE Math Misconceptions and Errors dataset card (README)", "Markdown",
        "research", "mae"); n += 1

    shutil.copy(SRC / "mae" / "data.json", DATA / f"D{n:02d}_mae_misconceptions.json")
    add(f"D{n:02d}", f"D{n:02d}_mae_misconceptions.json", "MaE: 55 algebra misconceptions with 220 diagnostic examples",
        "JSON", "learner", "mae", notes="researcher-designed; real misconception catalogue"); n += 1

    frac = re.compile(r"fraction|half|halves|third|quarter|percent|%|ratio|\d/\d", re.I)
    rows = [json.loads(line) for line in open(SRC / "mathdial" / "test.jsonl")]
    picked = [r for r in rows if frac.search(r["question"])][:60]
    with open(DATA / f"D{n:02d}_mathdial_fraction_dialogues.jsonl", "w") as fh:
        for r in picked:
            fh.write(json.dumps(r, ensure_ascii=False) + "\n")
    add(f"D{n:02d}", f"D{n:02d}_mathdial_fraction_dialogues.jsonl",
        "MathDial: 60 tutoring dialogues on fraction, ratio and percent problems (test split)", "JSONL", "learner",
        "mathdial", notes="teachers human, students LLM-simulated"); n += 1

    with open(SRC / "mathdial" / "test.csv", newline="") as fh:
        rows = list(csv.DictReader(fh))
    sample = random.sample(rows, 150)
    cols = ["qid", "question", "student_incorrect_solution", "teacher_described_confusion", "self-correctness"]
    with open(DATA / f"D{n:02d}_mathdial_confusions.csv", "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=cols)
        w.writeheader()
        for r in sample:
            w.writerow({c: r[c] for c in cols})
    add(f"D{n:02d}", f"D{n:02d}_mathdial_confusions.csv",
        "MathDial: 150 incorrect student solutions with teacher-described confusion", "CSV", "learner", "mathdial",
        notes="seeded random sample of the test split"); n += 1

    transcripts = json.loads((SRC / "manifest.json").read_text())["talkmoves_transcripts"]
    for t in transcripts:
        slug = re.sub(r"[^a-z0-9]+", "_", t.lower().replace(".xlsx", "")).strip("_")
        fname = f"D{n:02d}_talkmoves_{slug}.xlsx"
        shutil.copy(SRC / "talkmoves" / "transcripts" / t, DATA / fname)
        add(f"D{n:02d}", fname, f"TalkMoves classroom transcript: {t.replace('.xlsx', '')}", "XLSX", "learner",
            "talkmoves", notes="real K-12 classroom talk, teacher and student turns"); n += 1
    return n


# ---------------------------------------------------------------- synthetic
def build_synthetic(n: int) -> int:
    from docx import Document as Docx
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

    GOLD.mkdir(exist_ok=True)
    fname = f"D{n:02d}_synthetic_student_explanations.txt"
    (DATA / fname).write_text("\n".join(t for t, _ in S.EXPLANATIONS) + "\n", encoding="utf-8")
    with open(GOLD / "synthetic_labels.csv", "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["doc_id", "line", "text", "malrule_id", "malrule"])
        for i, (t, m) in enumerate(S.EXPLANATIONS, 1):
            w.writerow([f"D{n:02d}", i, t, m, S.MALRULES[m]])
    add(f"D{n:02d}", fname, "Student explanations of fraction, decimal and percent procedures", "TXT", "synthetic",
        "synthetic", True, notes="one line per named malrule; labels in gold/synthetic_labels.csv"); n += 1

    fname = f"D{n:02d}_synthetic_student_questions.txt"
    (DATA / fname).write_text("\n".join(S.QUESTIONS) + "\n", encoding="utf-8")
    add(f"D{n:02d}", fname, "Questions students ask during a fractions lesson", "TXT", "synthetic", "synthetic", True); n += 1

    d = Docx()
    d.add_heading(S.TEACHER_NOTES["title"], 0)
    for m in S.TEACHER_NOTES["meta"]:
        d.add_paragraph(m)
    d.add_heading("Observations", 1)
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Date", "Observation"
    for date, obs in S.TEACHER_NOTES["entries"]:
        row = table.add_row().cells
        row[0].text, row[1].text = date, obs
    d.add_heading("Next steps", 1)
    for s in S.TEACHER_NOTES["next_steps"]:
        d.add_paragraph(s, style="List Bullet")
    fname = f"D{n:02d}_synthetic_teacher_observation_notes.docx"
    d.save(DATA / fname)
    add(f"D{n:02d}", fname, "Teacher observation notes on class misconceptions", "DOCX", "synthetic", "synthetic", True); n += 1

    fname = f"D{n:02d}_synthetic_tutor_chat_log.json"
    (DATA / fname).write_text(json.dumps(S.TUTOR_CHAT, indent=2, ensure_ascii=False), encoding="utf-8")
    add(f"D{n:02d}", fname, "Tutor chat log with one learner over three sessions", "JSON", "synthetic", "synthetic", True,
        notes="longitudinal: misconception in session 1, corrected by session 3"); n += 1

    fname = f"D{n:02d}_synthetic_class_webpage.html"
    (DATA / fname).write_text(S.CLASS_PAGE_HTML, encoding="utf-8")
    add(f"D{n:02d}", fname, "Class web page: fractions unit lesson plan and schedule", "HTML", "synthetic", "synthetic", True,
        notes="includes navigation, cookie banner and footer boilerplate on purpose"); n += 1

    fname = f"D{n:02d}_synthetic_worksheet.pdf"
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(DATA / fname), pagesize=A4, title=S.WORKSHEET["title"])
    story = [Paragraph(S.WORKSHEET["title"], styles["Title"]), Paragraph(S.WORKSHEET["header"], styles["Normal"]),
             Spacer(1, 12),
             ListFlowable([ListItem(Paragraph(p, styles["Normal"])) for p in S.WORKSHEET["problems"]],
                          bulletType="1")]
    doc.build(story)
    add(f"D{n:02d}", fname, "Word-problem worksheet: fractions, decimals and percents", "PDF", "synthetic", "synthetic", True); n += 1
    return n


# ---------------------------------------------------------------- scale tier
def build_scale() -> None:
    (SCALE / "mathdial").mkdir(parents=True, exist_ok=True)
    (SCALE / "talkmoves").mkdir(parents=True, exist_ok=True)
    for f in ["train.jsonl", "test.jsonl"]:
        shutil.copy(SRC / "mathdial" / f, SCALE / "mathdial" / f)
    for f in ["train_data_504.xlsx", "test_data_63.xlsx"]:
        shutil.copy(SRC / "talkmoves" / f, SCALE / "talkmoves" / f)


def main() -> None:
    if DATA.exists():
        shutil.rmtree(DATA)
    DATA.mkdir()
    n = build_openstax(1)
    n = build_acl(n)
    n = build_dataset_docs(n)
    n = build_synthetic(n)
    build_scale()
    with open(DATA.parent / "results" / "document_catalog.csv", "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(catalog[0].keys()))
        w.writeheader()
        w.writerows(catalog)
    print(f"built {len(catalog)} documents")
    for c in catalog:
        print(f"  {c['doc_id']}  {c['format']:<22} {c['layer']:<9} {c['filename']}  {c['notes']}")


if __name__ == "__main__":
    main()
